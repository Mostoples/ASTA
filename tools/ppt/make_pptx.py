"""Menyusun PPTX dari keluaran build.js.

    python tools/ppt/make_pptx.py

Setiap slide = gambar latar (kartu, bayangan neumorphism, foto) + kotak teks asli
yang tetap bisa diedit di PowerPoint + video tertanam + transisi + catatan pembicara.
Hasil: "PRESENTASI FINAL HRIE 2026/KRPS001_Naurora Dewi Andrianto.pptx" dan .pdf
"""
import io, os, re, json, shutil, copy
from PIL import Image
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
DECK = os.path.join(HERE, 'out', 'deck')
DEST = os.path.join(ROOT, 'PRESENTASI FINAL HRIE 2026')
NAME = 'KRPS001_Naurora Dewi Andrianto'

PX = 6350                      # 1920 px = 13,333 in = 12.192.000 EMU
PT = 0.5                       # 1 px = 0,5 pt pada ukuran slide ini


def emu(v):
    return Emu(int(round(v * PX)))


def rgb(css):
    m = re.match(r'rgba?\((\d+),\s*(\d+),\s*(\d+)(?:,\s*([\d.]+))?', css)
    if not m:
        return None, 1.0
    a = float(m.group(4)) if m.group(4) is not None else 1.0
    return RGBColor(int(m.group(1)), int(m.group(2)), int(m.group(3))), a


def font_for(weight):
    """Keluarga Segoe UI yang dipakai Chromium untuk bobot CSS tersebut."""
    if weight >= 800:
        return 'Segoe UI Black', False
    if weight >= 700:
        return 'Segoe UI', True
    if weight >= 600:
        return 'Segoe UI Semibold', False
    return 'Segoe UI', False


def add_text(slide, t):
    one_line = t['h'] <= t['lineHeight'] * 1.4
    pad = 0.015 * t["w"] + 6
    x, w = t['x'], t['w'] + pad
    if t['align'] == 'center':
        x -= pad / 2
    elif t['align'] == 'right':
        x -= pad
    box = slide.shapes.add_textbox(emu(x), emu(t['y']), emu(w), emu(max(t['h'], t['lineHeight'])))
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = not one_line
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, runs in enumerate(t['paras']):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}.get(t['align'], PP_ALIGN.LEFT)
        p.line_spacing = Pt(t['lineHeight'] * PT)
        for r in runs:
            if not r['text']:
                continue
            run = p.add_run()
            run.text = r['text']
            f = run.font
            name, bold = font_for(r['weight'])
            f.name, f.bold, f.italic = name, bold, r['italic']
            f.size = Pt(round(r['size'] * PT * 2) / 2)
            col, alpha = rgb(r['color'])
            if col is not None:
                f.color.rgb = col
                if alpha < 0.99:
                    sf = f.color._xFill if hasattr(f.color, '_xFill') else None
                    srgb = run._r.find('.//' + qn('a:srgbClr'))
                    if srgb is not None:
                        a = etree.SubElement(srgb, qn('a:alpha'))
                        a.set('val', str(int(alpha * 100000)))
            if r['spacing']:
                run._r.get_or_add_rPr().set('spc', str(int(round(r['spacing'] * PT * 100))))
            # nama font juga untuk teks Asia/simbol, supaya tidak jatuh ke font tema
            rpr = run._r.get_or_add_rPr()
            for tag in ('a:ea', 'a:cs'):
                el = rpr.find(qn(tag))
                if el is None:
                    el = etree.SubElement(rpr, qn(tag))
                el.set('typeface', name)
    return box


def add_transition(slide, kind):
    sld = slide._element
    P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
    tr = etree.Element('{%s}transition' % P, spd='med')
    etree.SubElement(tr, '{%s}%s' % (P, 'push' if kind == 'push' else 'fade'))
    # urutan skema: cSld, clrMapOvr, transition, timing
    anchor = sld.find(qn('p:clrMapOvr'))
    if anchor is None:
        anchor = sld.find(qn('p:cSld'))
    anchor.addnext(tr)


def main():
    data = json.load(io.open(os.path.join(DECK, 'deck.json'), encoding='utf-8'))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(1920 * PX), Emu(1080 * PX)
    blank = prs.slide_layouts[6]
    tmp = os.path.join(DECK, '_jpg')
    os.makedirs(tmp, exist_ok=True)

    for d in data:
        s = prs.slides.add_slide(blank)
        bg = os.path.join(tmp, d['id'] + '.jpg')
        Image.open(os.path.join(DECK, d['id'] + '-bg.png')).convert('RGB').save(bg, quality=92, subsampling=0)
        s.shapes.add_picture(bg, 0, 0, prs.slide_width, prs.slide_height)
        for t in d['texts']:
            add_text(s, t)
        for v in d['videos']:
            s.shapes.add_movie(os.path.join(HERE, v['src']), emu(v['x']), emu(v['y']), emu(v['w']), emu(v['h']),
                               poster_frame_image=os.path.join(HERE, v['poster']), mime_type='video/mp4')
        if d['notes']:
            s.notes_slide.notes_text_frame.text = d['notes']
        add_transition(s, d['transition'])

    prs.core_properties.title = 'ASTA — Presentasi Final HRIE 2026'
    prs.core_properties.author = 'Tim ASTA · SMA Negeri 1 Surakarta'
    os.makedirs(DEST, exist_ok=True)
    out = os.path.join(DEST, NAME + '.pptx')
    prs.save(out)
    shutil.copyfile(os.path.join(DECK, 'deck.pdf'), os.path.join(DEST, NAME + '.pdf'))
    print(out, '%.1f MB' % (os.path.getsize(out) / 1e6))


if __name__ == '__main__':
    main()
